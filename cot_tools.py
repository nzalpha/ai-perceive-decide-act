from mcp.server.fastmcp import FastMCP
from mcp.types import TextContent
from rich.console import Console
from rich.panel import Panel
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
import math

console = Console()
mcp = FastMCP("PowerPointASCIICalculator")

@mcp.tool()
def get_ascii_values(text: str) -> TextContent:
    """Get ASCII values for each character in the given text"""
    console.print("[blue]FUNCTION CALL:[/blue] get_ascii_values()")
    console.print(f"[blue]Text:[/blue] {text}")
    
    ascii_values = [ord(char) for char in text]
    result = {char: ord(char) for char in text}
    
    console.print(Panel(
        "\n".join([f"{char}: {val}" for char, val in result.items()]),
        title="ASCII Values",
        border_style="cyan"
    ))
    
    return TextContent(
        type="text",
        text=str(ascii_values)
    )

@mcp.tool()
def calculate_exponential_sum(ascii_values: list) -> TextContent:
    """Calculate the sum of exponentials (e^x) for each ASCII value"""
    console.print("[blue]FUNCTION CALL:[/blue] calculate_exponential_sum()")
    console.print(f"[blue]ASCII Values:[/blue] {ascii_values}")
    
    try:
        exponentials = [math.exp(val) for val in ascii_values]
        total_sum = sum(exponentials)
        
        console.print(Panel(
            f"e^{ascii_values[0]} + e^{ascii_values[1]} + ... = {total_sum}",
            title="Exponential Sum",
            border_style="green"
        ))
        
        return TextContent(
            type="text",
            text=str(total_sum)
        )
    except Exception as e:
        console.print(f"[red]Error:[/red] {str(e)}")
        return TextContent(
            type="text",
            text=f"Error: {str(e)}"
        )

@mcp.tool()
def show_reasoning(steps: list) -> TextContent:
    """Display the step-by-step reasoning process"""
    console.print("[blue]FUNCTION CALL:[/blue] show_reasoning()")
    
    for i, step in enumerate(steps, 1):
        console.print(Panel(
            f"{step}",
            title=f"Step {i}",
            border_style="cyan"
        ))
    
    return TextContent(
        type="text",
        text="Steps displayed"
    )

@mcp.tool()
def add_rectangle_to_ppt(file_path: str, text: str, output_path: str = None) -> TextContent:
    """Open PowerPoint file, add a rectangle with text, and save it"""
    console.print("[blue]FUNCTION CALL:[/blue] add_rectangle_to_ppt()")
    console.print(f"[blue]Input File:[/blue] {file_path}")
    console.print(f"[blue]Text:[/blue] {text}")
    
    try:
        import os
        # Open the presentation if it exists, otherwise create a new one
        if os.path.exists(file_path):
            prs = Presentation(file_path)
            console.print(f"[green]Opened existing file:[/green] {file_path}")
        else:
            prs = Presentation()
            console.print(f"[yellow]Creating new presentation (file not found):[/yellow] {file_path}")
        
        # Get the first slide (or create one if none exists)
        if len(prs.slides) == 0:
            # Add a blank slide
            blank_slide_layout = prs.slide_layouts[6]  # Blank layout
            slide = prs.slides.add_slide(blank_slide_layout)
        else:
            slide = prs.slides[0]
        
        # Add a rectangle shape
        left = Inches(2)
        top = Inches(3)
        width = Inches(5)
        height = Inches(1.5)
        
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left, top, width, height
        )
        
        # Style the rectangle
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(68, 114, 196)  # Blue color
        shape.line.color.rgb = RGBColor(0, 0, 0)  # Black border
        
        # Add text to the rectangle
        text_frame = shape.text_frame
        text_frame.clear()
        p = text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)  # White text
        
        # Center align the text
        from pptx.enum.text import PP_ALIGN
        p.alignment = PP_ALIGN.CENTER
        
        # Save the presentation
        if output_path is None:
            output_path = "output.pptx"
        
        prs.save(output_path)
        
        console.print(Panel(
            f"Rectangle added successfully!\nSaved to: {output_path}",
            title="Success",
            border_style="green"
        ))
        
        return TextContent(
            type="text",
            text=f"Success: Rectangle added and saved to {output_path}"
        )
    except Exception as e:
        console.print(f"[red]Error:[/red] {str(e)}")
        return TextContent(
            type="text",
            text=f"Error: {str(e)}"
        )

@mcp.tool()
def verify_calculation(description: str, expected_value: float, actual_value: float) -> TextContent:
    """Verify if a calculation matches the expected value"""
    console.print("[blue]FUNCTION CALL:[/blue] verify_calculation()")
    console.print(f"[blue]Verifying:[/blue] {description}")
    
    is_correct = abs(actual_value - expected_value) < 1e-6
    
    if is_correct:
        console.print(f"[green]✓ Correct! {description}: {actual_value}[/green]")
    else:
        console.print(f"[red]✗ Incorrect! Expected {expected_value}, got {actual_value}[/red]")
    
    return TextContent(
        type="text",
        text=str(is_correct)
    )

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == "dev":
        mcp.run()
    else:
        mcp.run(transport="stdio")




###########################################################
########################################################### Tools for the agent ###########################################################
# from mcp.server.fastmcp import FastMCP
# from mcp.types import TextContent
# from rich.console import Console
# from rich.panel import Panel
# from pptx import Presentation
# from pptx.util import Inches, Pt
# from pptx.enum.shapes import MSO_SHAPE
# from pptx.dml.color import RGBColor
# import math
# from models import AsciiValueInput, OperationValueOutput, ExponentialSumInput, ReasoningStepsInputs

# console = Console()
# mcp = FastMCP("PowerPointASCIICalculator")

# # @mcp.tool()
# # def get_ascii_values(text: str) -> TextContent:
# #     """Get ASCII values for each character in the given text"""
# #     console.print("[blue]FUNCTION CALL:[/blue] get_ascii_values()")
# #     console.print(f"[blue]Text:[/blue] {text}")
    
# #     ascii_values = [ord(char) for char in text]
# #     result = {char: ord(char) for char in text}
    
# #     console.print(Panel(
# #         "\n".join([f"{char}: {val}" for char, val in result.items()]),
# #         title="ASCII Values",
# #         border_style="cyan"
# #     ))
    
# #     return TextContent(
# #         type="text",
# #         text=str(ascii_values)
# #     )


# @mcp.tool()
# def get_ascii_values(text: str) -> OperationValueOutput:
#     """Get ASCII values for each character in the given text"""
#     input = AsciiValueInput(text=text) 
#     console.print("[blue]FUNCTION CALL:[/blue] get_ascii_values()")
#     console.print(f"[blue]Text:[/blue] {input.text}")

    
#     ascii_values = [ord(char) for char in input.text]
#     result = {char: ord(char) for char in input.text}
    
#     console.print(Panel(
#         "\n".join([f"{char}: {val}" for char, val in result.items()]),
#         title="ASCII Values",
#         border_style="cyan"
#     ))
    
#     return OperationValueOutput(
#         type="text",
#         text=str(ascii_values)
#     )

# @mcp.tool()
# def calculate_exponential_sum(ascii_values: list) -> OperationValueOutput:
#     """Calculate the sum of exponentials (e^x) for each ASCII value"""
#     input = ExponentialSumInput(ascii_values=ascii_values)
#     console.print("[blue]FUNCTION CALL:[/blue] calculate_exponential_sum()")
#     console.print(f"[blue]ASCII Values:[/blue] {input.ascii_values}")
    
#     try:
#         exponentials = [math.exp(val) for val in input.ascii_values]
#         total_sum = sum(exponentials)
        
#         console.print(Panel(
#             f"e^{ascii_values[0]} + e^{ascii_values[1]} + ... = {total_sum}",
#             title="Exponential Sum",
#             border_style="green"
#         ))
        
#         return OperationValueOutput(
#             type="text",
#             text=str(total_sum)
#         )
#     except Exception as e:
#         console.print(f"[red]Error:[/red] {str(e)}")
#         return OperationValueOutput(
#             type="text",
#             text=f"Error: {str(e)}"
#         )

# @mcp.tool()
# def show_reasoning(steps: list) -> OperationValueOutput:
#     """Display the step-by-step reasoning process"""
#     input = ReasoningStepsInputs(steps=steps)
#     console.print("[blue]FUNCTION CALL:[/blue] show_reasoning()")
    
#     for i, step in enumerate(input.steps, 1):
#         console.print(Panel(
#             f"{step}",
#             title=f"Step {i}",
#             border_style="cyan"
#         ))
    
#     return OperationValueOutput(
#         type="text",
#         text="Steps displayed"
#     )

# @mcp.tool()
# def add_rectangle_to_ppt(file_path: str, text: str, output_path: str = None) -> TextContent:
#     """Open PowerPoint file, add a rectangle with text, and save it"""
#     console.print("[blue]FUNCTION CALL:[/blue] add_rectangle_to_ppt()")
#     console.print(f"[blue]Input File:[/blue] {file_path}")
#     console.print(f"[blue]Text:[/blue] {text}")
    
#     try:
#         import os
#         # Open the presentation if it exists, otherwise create a new one
#         if os.path.exists(file_path):
#             prs = Presentation(file_path)
#             console.print(f"[green]Opened existing file:[/green] {file_path}")
#         else:
#             prs = Presentation()
#             console.print(f"[yellow]Creating new presentation (file not found):[/yellow] {file_path}")
        
#         # Get the first slide (or create one if none exists)
#         if len(prs.slides) == 0:
#             # Add a blank slide
#             blank_slide_layout = prs.slide_layouts[6]  # Blank layout
#             slide = prs.slides.add_slide(blank_slide_layout)
#         else:
#             slide = prs.slides[0]
        
#         # Add a rectangle shape
#         left = Inches(2)
#         top = Inches(3)
#         width = Inches(5)
#         height = Inches(1.5)
        
#         shape = slide.shapes.add_shape(
#             MSO_SHAPE.RECTANGLE,
#             left, top, width, height
#         )
        
#         # Style the rectangle
#         shape.fill.solid()
#         shape.fill.fore_color.rgb = RGBColor(68, 114, 196)  # Blue color
#         shape.line.color.rgb = RGBColor(0, 0, 0)  # Black border
        
#         # Add text to the rectangle
#         text_frame = shape.text_frame
#         text_frame.clear()
#         p = text_frame.paragraphs[0]
#         p.text = text
#         p.font.size = Pt(24)
#         p.font.bold = True
#         p.font.color.rgb = RGBColor(255, 255, 255)  # White text
        
#         # Center align the text
#         from pptx.enum.text import PP_ALIGN
#         p.alignment = PP_ALIGN.CENTER
        
#         # Save the presentation
#         if output_path is None:
#             output_path = "output.pptx"
        
#         prs.save(output_path)
        
#         console.print(Panel(
#             f"Rectangle added successfully!\nSaved to: {output_path}",
#             title="Success",
#             border_style="green"
#         ))
        
#         return TextContent(
#             type="text",
#             text=f"Success: Rectangle added and saved to {output_path}"
#         )
#     except Exception as e:
#         console.print(f"[red]Error:[/red] {str(e)}")
#         return TextContent(
#             type="text",
#             text=f"Error: {str(e)}"
#         )

# @mcp.tool()
# def verify_calculation(description: str, expected_value: float, actual_value: float) -> TextContent:
#     """Verify if a calculation matches the expected value"""
#     console.print("[blue]FUNCTION CALL:[/blue] verify_calculation()")
#     console.print(f"[blue]Verifying:[/blue] {description}")
    
#     is_correct = abs(actual_value - expected_value) < 1e-6
    
#     if is_correct:
#         console.print(f"[green]✓ Correct! {description}: {actual_value}[/green]")
#     else:
#         console.print(f"[red]✗ Incorrect! Expected {expected_value}, got {actual_value}[/red]")
    
#     return TextContent(
#         type="text",
#         text=str(is_correct)
#     )

# if __name__ == "__main__":
#     import sys
#     if len(sys.argv) > 1 and sys.argv[1] == "dev":
#         mcp.run()
#     else:
#         mcp.run(transport="stdio")







# # from mcp.server.fastmcp import FastMCP
# # from mcp.types import TextContent
# # from rich.console import Console
# # from rich.panel import Panel

# # console = Console()
# # mcp = FastMCP("CoTCalculator")

# # @mcp.tool()
# # def show_reasoning(steps: list) -> TextContent:
# #     """Show the step-by-step reasoning process"""
# #     console.print("[blue]FUNCTION CALL:[/blue] show_reasoning()")
# #     for i, step in enumerate(steps, 1):
# #         console.print(Panel(
# #             f"{step}",
# #             title=f"Step {i}",
# #             border_style="cyan"
# #         ))
# #     return TextContent(
# #         type="text",
# #         text="Reasoning shown"
# #     )

# # @mcp.tool()
# # def calculate(expression: str) -> TextContent:
# #     """Calculate the result of an expression"""
# #     console.print("[blue]FUNCTION CALL:[/blue] calculate()")
# #     console.print(f"[blue]Expression:[/blue] {expression}")
# #     try:
# #         result = eval(expression)
# #         console.print(f"[green]Result:[/green] {result}")
# #         return TextContent(
# #             type="text",
# #             text=str(result)
# #         )
# #     except Exception as e:
# #         console.print(f"[red]Error:[/red] {str(e)}")
# #         return TextContent(
# #             type="text",
# #             text=f"Error: {str(e)}"
# #         )

# # @mcp.tool()
# # def verify(expression: str, expected: float) -> TextContent:
# #     """Verify if a calculation is correct"""
# #     console.print("[blue]FUNCTION CALL:[/blue] verify()")
# #     console.print(f"[blue]Verifying:[/blue] {expression} = {expected}")
# #     try:
# #         actual = float(eval(expression))
# #         is_correct = abs(actual - float(expected)) < 1e-10
        
# #         if is_correct:
# #             console.print(f"[green]✓ Correct! {expression} = {expected}[/green]")
# #         else:
# #             console.print(f"[red]✗ Incorrect! {expression} should be {actual}, got {expected}[/red]")
            
# #         return TextContent(
# #             type="text",
# #             text=str(is_correct)
# #         )
# #     except Exception as e:
# #         console.print(f"[red]Error:[/red] {str(e)}")
# #         return TextContent(
# #             type="text",
# #             text=f"Error: {str(e)}"
# #         )

# # if __name__ == "__main__":
# #     import sys
# #     if len(sys.argv) > 1 and sys.argv[1] == "dev":
# #         mcp.run()
# #     else:
# #         mcp.run(transport="stdio")
